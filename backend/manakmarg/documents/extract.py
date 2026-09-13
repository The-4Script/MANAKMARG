"""Text and table extraction with page numbers for uploaded PDF, DOCX and TXT files.

Plain-text tables written with ``|`` separators are read as table rows so simple reports work without a PDF.
"""

from dataclasses import dataclass, field
from pathlib import Path

from manakmarg.normalize.text import clean_ws


@dataclass(frozen=True)
class PageText:
    number: int
    lines: tuple[str, ...]


@dataclass(frozen=True)
class TableRows:
    page: int
    rows: tuple[tuple[str, ...], ...]


@dataclass
class ExtractedDoc:
    filename: str
    pages: list[PageText] = field(default_factory=list)
    tables: list[TableRows] = field(default_factory=list)

    @property
    def text_line_count(self) -> int:
        return sum(len(page.lines) for page in self.pages)


def _split_text(page_number: int, text: str, doc: ExtractedDoc) -> None:
    lines: list[str] = []
    table: list[tuple[str, ...]] = []
    for raw in text.splitlines():
        if raw.count("|") >= 2 or (raw.strip().startswith("|") and raw.count("|") >= 1):
            cells = tuple(clean_ws(cell) for cell in raw.strip().strip("|").split("|"))
            if not all(set(cell) <= set("-: ") for cell in cells):
                table.append(cells)
            continue
        if table:
            doc.tables.append(TableRows(page_number, tuple(table)))
            table = []
        line = clean_ws(raw)
        if line:
            lines.append(line)
    if table:
        doc.tables.append(TableRows(page_number, tuple(table)))
    doc.pages.append(PageText(page_number, tuple(lines)))


def extract_document(path: Path, filename: str | None = None) -> ExtractedDoc:
    path = Path(path)
    doc = ExtractedDoc(filename or path.name)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from manakmarg.ingest.documents import pdf_pages

        for page in pdf_pages(path):
            _split_text(page["page"], page["text"], doc)
            for table in page["tables"]:
                rows = tuple(tuple(clean_ws(cell or "") for cell in row) for row in table)
                if rows:
                    doc.tables.append(TableRows(page["page"], rows))
    elif suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        _split_text(1, "\n".join(paragraph.text for paragraph in document.paragraphs), doc)
        for table in document.tables:
            rows = tuple(tuple(clean_ws(cell.text) for cell in row.cells) for row in table.rows)
            if rows:
                doc.tables.append(TableRows(1, rows))
    else:
        _split_text(1, path.read_text(encoding="utf-8", errors="replace"), doc)
    return doc

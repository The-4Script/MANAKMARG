"""Data inventory of the supplied workbooks (plan Task 2.4).

Profiles every sheet of every workbook in the data folder — structure, headers, row counts, duplicates,
value gaps, classification, source and authority — and compares the exports with each other using the
canonical designation key. Written to ``data/manifests/data_inventory.json`` and ``docs/DATA_INVENTORY.md``.
"""

import collections
import json
import re
import zipfile
from pathlib import Path

import openpyxl

from manakmarg.ingest import sources
from manakmarg.ingest.excel_standards import (
    SOURCE_MINISTRY,
    SOURCE_SECOND,
    SOURCE_TOTAL,
    ExportFile,
    classify_export,
    read_export,
)
from manakmarg.normalize.is_number import parse_designation
from manakmarg.normalize.text import clean_ws

_SOURCE_BY_CLASSIFICATION = {
    "total_export": SOURCE_TOTAL,
    "untitled_export": SOURCE_SECOND,
    "ministry_node": SOURCE_MINISTRY,
}

_DESCRIPTIONS = {
    "total_export": "Full list of published Indian Standards exported from the BIS Standards portal (title cell 'Total').",
    "untitled_export": "Second full export of published Indian Standards with an empty title cell; repeated rows point to a "
    "view with multi-membership such as Group-wise.",
    "ministry_node": "Published standards classified under one ministry or ministry-department node named in the title cell.",
    "unknown_node": "Export whose title cell is not a recognised classification; not ingested until its dimension is confirmed.",
}

_RELATIONSHIPS = {
    "total_export": "Master list; every other export is compared with it by canonical designation key.",
    "untitled_export": "Overlaps the master list almost entirely; contributes designations absent from it and the "
    "'present in second export' flag.",
    "ministry_node": "Links standards to a ministry node. 'Ministry – Department' nodes are not subsets of their parent "
    "ministry's list.",
    "unknown_node": "Not established.",
}

_NOTES = {
    "total_export": "No Department or Group column: classification is not derivable from this file.",
    "untitled_export": "The export does not name its view; the Group-wise interpretation is an inference from its row pattern.",
    "ministry_node": "Ministry exports supplied cover only part of the ministry list.",
    "unknown_node": "Reported and skipped by ingestion.",
}

_IMPORTANT_COLUMNS = ["Standard Number", "Title", "Date of Publish", "Type of Standard", "Degree of Equivalence"]


def _sheet_states(path: Path) -> dict[str, str]:
    try:
        with zipfile.ZipFile(path) as archive:
            workbook_xml = archive.read("xl/workbook.xml").decode("utf-8", errors="replace")
    except (KeyError, zipfile.BadZipFile):
        return {}
    states = {}
    for tag in re.findall(r"<sheet\b[^>]*>", workbook_xml):
        name = re.search(r'name="([^"]*)"', tag)
        state = re.search(r'state="([^"]*)"', tag)
        if name:
            states[name.group(1)] = state.group(1) if state else "visible"
    return states


def _sheet_overviews(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        overviews = []
        for sheet in workbook.worksheets:
            header_row, headers = None, []
            for index, row in enumerate(sheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
                values = [clean_ws(str(cell)) if cell is not None else "" for cell in row]
                if "Standard Number" in values and "Title" in values:
                    header_row, headers = index, [value for value in values if value]
                    break
            overviews.append(
                {
                    "sheet": sheet.title,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "header_row": header_row,
                    "headers": headers,
                }
            )
        return overviews
    finally:
        workbook.close()


def _profile_export(export: ExportFile) -> tuple[dict, set[str]]:
    classification = classify_export(export)
    raw_values = [row.designation_raw for row in export.rows]
    parsed = {raw: parse_designation(raw) for raw in set(raw_values)}
    keys = {designation.std_key for designation in parsed.values() if designation is not None}

    raw_counts = collections.Counter(raw_values)
    repeated = {raw: count for raw, count in raw_counts.items() if count > 1}
    raws_by_key = collections.defaultdict(set)
    for raw, designation in parsed.items():
        if designation is not None:
            raws_by_key[designation.std_key].add(raw)

    source_id = _SOURCE_BY_CLASSIFICATION.get(classification)
    definition = sources.REGISTRY.get(source_id) if source_id else None
    profile = {
        "row_count": len(export.rows),
        "title_cell": export.title_cell,
        "generated_on": export.generated_on.isoformat() if export.generated_on else None,
        "classification": classification,
        "domain": "Indian Standards metadata"
        + (" — ministry classification" if classification == "ministry_node" else ""),
        "description": _DESCRIPTIONS[classification],
        "important_columns": _IMPORTANT_COLUMNS,
        "candidate_key": "Standard Number, normalised to std_key (prefix + number + part/section + suffix + year)",
        "duplicate_characteristics": {
            "repeated_designations": sum(count - 1 for count in repeated.values()),
            "sample_repeated": sorted(repeated)[:5],
            "raw_variants_collapsing_to_one_key": sum(1 for raws in raws_by_key.values() if len(raws) > 1),
        },
        "relationship_to_other_data": _RELATIONSHIPS[classification],
        "data_quality_issues": {
            "null_publish_dates": sum(1 for row in export.rows if row.publish_date_raw is None),
            "type_unspecified": sum(1 for row in export.rows if row.standard_type is None),
            "equivalence_unspecified": sum(1 for row in export.rows if row.degree is None),
            "missing_titles": sum(1 for row in export.rows if not row.title),
            "unparseable_designations": sum(1 for designation in parsed.values() if designation is None),
            "designations_with_format_flags": sum(
                1 for designation in parsed.values() if designation is not None and designation.flags
            ),
        },
        "source_type": "excel_export",
        "source_id": source_id,
        "authority_level": definition.authority if definition else "unknown",
        "notes": _NOTES[classification],
        "sample_records": [
            {
                "sheet_row": row.sheet_row,
                "standard_number": row.designation_raw,
                "date_of_publish": row.publish_date_raw,
                "title": row.title,
                "type_of_standard": row.standard_type,
                "degree_of_equivalence": row.degree,
            }
            for row in export.rows[:3]
        ],
    }
    return profile, keys


def build_inventory(data_dir: Path) -> dict:
    data_dir = Path(data_dir)
    workbooks = []
    keys_by_classification: dict[str, list[tuple[str, set[str]]]] = collections.defaultdict(list)

    for path in sorted(data_dir.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        states = _sheet_states(path)
        overviews = _sheet_overviews(path)
        export = read_export(path)
        profile, keys = _profile_export(export)
        keys_by_classification[profile["classification"]].append((path.name, keys))

        sheets = []
        for overview in overviews:
            entry = {
                "sheet": overview["sheet"],
                "state": states.get(overview["sheet"], "visible"),
                "max_row": overview["max_row"],
                "max_column": overview["max_column"],
                "header_row": overview["header_row"],
                "headers": overview["headers"],
                "column_count": len(overview["headers"]),
            }
            if overview["sheet"] == export.sheet_name:
                entry.update(profile)
            sheets.append(entry)
        workbooks.append({"file": path.name, "bytes": path.stat().st_size, "sheets": sheets})

    total = set().union(*(keys for _, keys in keys_by_classification["total_export"]))
    untitled = set().union(*(keys for _, keys in keys_by_classification["untitled_export"]))
    ministry_counts = collections.Counter(
        key for _, keys in keys_by_classification["ministry_node"] for key in keys
    )
    not_in_full_exports = sorted(set(ministry_counts) - (total | untitled))

    cross_file = {
        "total_vs_untitled": (
            {
                "shared": len(total & untitled),
                "only_in_total": len(total - untitled),
                "only_in_untitled": len(untitled - total),
            }
            if total and untitled
            else None
        ),
        "ministry_exports": {
            "files": len(keys_by_classification["ministry_node"]),
            "distinct_standards": len(ministry_counts),
            "standards_in_more_than_one_node": sum(1 for count in ministry_counts.values() if count > 1),
            "not_in_total_or_untitled": len(not_in_full_exports),
            "examples_not_in_total_or_untitled": not_in_full_exports[:10],
        },
    }
    known_gaps = [
        "None of the supplied exports has a per-row Department or Group column, so BIS Department/Group "
        "classification is not available.",
    ]
    if keys_by_classification["ministry_node"]:
        known_gaps.append(
            f"Ministry-wise classification is partial: {len(keys_by_classification['ministry_node'])} ministry "
            "node export(s) supplied."
        )
    return {
        "generated_by": "manakmarg.ingest.inventory",
        "data_dir": data_dir.name,
        "workbook_count": len(workbooks),
        "workbooks": workbooks,
        "cross_file": cross_file,
        "known_gaps": known_gaps,
    }


def write_inventory(inventory: dict, json_path: Path, markdown_path: Path) -> None:
    json_path = Path(json_path)
    markdown_path = Path(markdown_path)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(inventory, indent=2, ensure_ascii=False), encoding="utf-8")

    lines = [
        "# Data inventory — supplied workbooks",
        "",
        f"Generated by `{inventory['generated_by']}` from the `{inventory['data_dir']}/` folder. "
        "The workbooks themselves are never modified.",
        "",
        "| File | Sheet | State | Classification | Title cell | Rows | Columns | Generated on | Source | Authority |",
        "|---|---|---|---|---|---:|---:|---|---|---|",
    ]
    for workbook in inventory["workbooks"]:
        for sheet in workbook["sheets"]:
            lines.append(
                f"| {workbook['file']} | {sheet['sheet']} | {sheet['state']} | {sheet.get('classification', '')} | "
                f"{sheet.get('title_cell') or '(empty)'} | {sheet.get('row_count', '')} | {sheet['column_count']} | "
                f"{sheet.get('generated_on') or ''} | {sheet.get('source_id') or ''} | {sheet.get('authority_level', '')} |"
            )

    lines += ["", "## Per-file quality and duplicates", ""]
    lines += [
        "| File | Null publish dates | Type '-' | Equivalence '-' | Repeated designations | Format-flagged designations | Unparseable |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for workbook in inventory["workbooks"]:
        for sheet in workbook["sheets"]:
            issues = sheet.get("data_quality_issues")
            if not issues:
                continue
            duplicates = sheet["duplicate_characteristics"]
            lines.append(
                f"| {workbook['file']} | {issues['null_publish_dates']} | {issues['type_unspecified']} | "
                f"{issues['equivalence_unspecified']} | {duplicates['repeated_designations']} | "
                f"{issues['designations_with_format_flags']} | {issues['unparseable_designations']} |"
            )

    cross = inventory["cross_file"]
    lines += ["", "## Relationships between exports", ""]
    if cross["total_vs_untitled"]:
        overlap = cross["total_vs_untitled"]
        lines.append(
            f"- Total vs untitled export: {overlap['shared']:,} shared designations, {overlap['only_in_total']:,} only "
            f"in the total export, {overlap['only_in_untitled']:,} only in the untitled export."
        )
    ministry = cross["ministry_exports"]
    lines.append(
        f"- Ministry exports: {ministry['files']} file(s), {ministry['distinct_standards']:,} distinct standards, "
        f"{ministry['standards_in_more_than_one_node']:,} under more than one node, "
        f"{ministry['not_in_total_or_untitled']:,} not present in either full export"
        + (f" (e.g. {', '.join(ministry['examples_not_in_total_or_untitled'][:5])})." if ministry["examples_not_in_total_or_untitled"] else ".")
    )
    lines += ["", "## Known gaps", ""]
    lines += [f"- {gap}" for gap in inventory["known_gaps"]]
    lines += ["", "## Column descriptions", ""]
    lines += [
        "- **Standard Number** — the published designation; normalised to a canonical key while the original text is kept.",
        "- **Date of Publish** — publication date as exported (`DD Mon YYYY`); blank for some older standards.",
        "- **Title** — standard title including revision and amendment markers.",
        "- **Type of Standard** — e.g. Product Specification, Methods of Tests; `-` means unspecified.",
        "- **Degree of Equivalence** — relation to international standards; `-` means unspecified.",
    ]
    markdown_path.write_text("\n".join(lines) + "\n", encoding="utf-8")

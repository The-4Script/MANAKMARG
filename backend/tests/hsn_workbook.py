"""Test helper: small HSN master workbooks written at test time (the real workbook is never touched by tests)."""

from pathlib import Path

import openpyxl

HSN_ROWS = [
    ("73", "ARTICLES OF IRON OR STEEL"),
    ("7305", "OTHER TUBES AND PIPES (FOR EXAMPLE, WELDED, RIVETED OR SIMILARLY CLOSED), HAVING CIRCULAR CROSS-SECTIONS"),
    ("730531", "OTHER, WELDED : LONGITUDINALLY WELDED"),
    ("73053121", "Non-galvanised, of iron : Clad, plated or coated"),
    ("73053129", "Non-galvanised, of iron : Other"),
    ("73053911", "Galvanised : Of iron"),
    ("74", "COPPER AND ARTICLES THEREOF"),
    ("7408", "COPPER WIRE"),
    ("74081110", "WIRE OF REFINED COPPER : OF WHICH THE MAXIMUM CROSS-SECTIONAL DIMENSION EXCEEDS 6 MM"),
    ("74081990", "COPPER WIRE OF REFINED COPPER OTHER"),
    ("7414", "OMITTED"),
    ("74142010", "CLOTH (INCLUDING ENDLESS BANDS), GRILL AND NETTING, OF COPPER WIRE; EXPANDED METAL OF COPPER - CLOTH : ENDLESS BANDS, FOR MACHINERY"),
    ("7323", "TABLE, KITCHEN OR OTHER HOUSEHOLD ARTICLES AND PARTS THEREOF, OF IRON OR STEEL"),
    ("73239310", "TABLE, KITCHEN OR OTHER HOUSEHOLD ARTICLES OF STAINLESS STEEL : UTENSILS"),
    ("2307 00", "WINE LEES; ARGOL"),
    ("0101", "LIVE HORSES, ASSES, MULES AND HINNIES."),
    ("85", "  ELECTRICAL MACHINERY  AND EQUIPMENT  "),
]


def write_hsn_workbook(directory: Path, *, name: str = "master_codes.xlsx", rows=None, extra_rows=True) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "HSN_MSTR"
    sheet.append(("HSN_CD", "HSN_Description"))
    for row in rows if rows is not None else HSN_ROWS:
        sheet.append(row)
    if extra_rows:
        sheet.append(("7408", "duplicate row must be skipped"))
        sheet.append(("12AB", "invalid code must be rejected"))
        sheet.append(("99", None))
        sheet.append((None, None))
    services = workbook.create_sheet("SAC_MSTR")
    services.append(("SAC_CD", "SAC_Description"))
    services.append((9954, "Construction services"))
    path = directory / name
    workbook.save(path)
    return path


def write_decoy_workbook(directory: Path) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    workbook.active.append(("Standard No.", "Title"))
    workbook.active.append(("IS 2062", "Structural steel"))
    path = directory / "aaa_not_hsn.xlsx"
    workbook.save(path)
    return path

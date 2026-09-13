"""Test helper: build small workbooks in the layout of the BIS 'Published Standards' exports."""

import openpyxl

HEADERS = ["Sl#", "Standard Number", "Date of Publish", "Title", "Type of Standard", "Degree of Equivalence"]


def make_export(path, title, rows, generated="Generated On:\nSep 12, 2026 08:17 PM"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Published Standards"
    sheet["C1"] = title
    sheet.merge_cells("C1:E1")
    sheet["F1"] = generated
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=2, column=column, value=header)
    for offset, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=offset, column=column, value=value)
    workbook.save(path)


def build_sample_exports(directory):
    """Three exports: a master list, an untitled second export and one ministry export."""
    directory.mkdir(parents=True, exist_ok=True)
    utensils = ("IS 14756:2024", "22 Nov 2024", "Stainless Steel Utensils", "Product Specification", "Indigenous")
    slag = ("IS 455:2015", "01 Jan 2016", "Portland Slag Cement", "Product Specification", "Indigenous")
    make_export(
        directory / "1.xlsx",
        "Total",
        [("1", *utensils), ("2", "IS 269 : 2015", None, "Ordinary Portland Cement", "-", "-"), ("3", *slag)],
    )
    door = ("IS 4020 (Part 8):1998", "31 Mar 1998", "Door shutters", "Methods of Tests", "Indigenous")
    make_export(
        directory / "2.xlsx",
        None,
        [("1", *utensils), ("2", *slag), ("3", *door), ("4", *door)],
        generated="Generated On:\nSep 12, 2026 08:32 PM",
    )
    make_export(
        directory / "File_Published_Standards_List_2026-09-12_203745.xlsx",
        "Ministry of Coal",
        [("1", *slag), ("2", "IS 7328:2020", "30 Jun 2020", "Polyethylene material", "Product Specification", "Indigenous")],
        generated="Generated On:\nSep 12, 2026 08:37 PM",
    )
    return directory

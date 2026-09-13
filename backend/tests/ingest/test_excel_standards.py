import json
from datetime import date, datetime

import openpyxl
import pytest
import sqlalchemy as sa

from manakmarg.core.clock import IST
from manakmarg.db import schema
from manakmarg.db.engine import get_engine, init_db
from manakmarg.ingest import sources
from manakmarg.ingest.excel_standards import classify_export, ingest_standard_exports, read_export

HEADERS = ["Sl#", "Standard Number", "Date of Publish", "Title", "Type of Standard", "Degree of Equivalence"]


def make_export(path, title, rows, generated="Generated On:\nSep 12, 2026 08:17 PM"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Published Standards"
    sheet["C1"] = title
    sheet.merge_cells("C1:E1")
    sheet["F1"] = generated
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(row=2, column=column, value=header)
    for offset, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=offset, column=column, value=value)
    workbook.save(path)


MASTER_ROWS = [
    ("1", "IS 14756:2024", "22 Nov 2024", "Stainless Steel Utensils — Specification ( Third Revision )", "Product Specification", "Indigenous"),
    ("2", "IS 269 : 2015", None, "Ordinary Portland Cement — Specification", "-", "-"),
    ("3", "IS/ISO 80000-9:2019", "18 May 2023", "Quantities and units Part 9: Physical chemistry", "Others", "Identical under single numbering"),
]


@pytest.fixture
def engine(tmp_path):
    eng = get_engine(tmp_path / "standards.sqlite3")
    init_db(eng)
    with eng.begin() as conn:
        sources.sync_registry(conn)
    yield eng
    eng.dispose()


@pytest.fixture
def data_dir(tmp_path):
    directory = tmp_path / "data"
    directory.mkdir()
    make_export(directory / "1.xlsx", "Total", MASTER_ROWS)
    make_export(
        directory / "2.xlsx",
        None,
        [
            ("1", "IS 14756:2024", "22 Nov 2024", "Stainless Steel Utensils — Specification ( Third Revision )", "Product Specification", "Indigenous"),
            ("2", "IS 4020 (Part 8):1998", "31 Mar 1998", "Door shutters - Methods of tests Part 8", "Methods of Tests", "Indigenous"),
            ("3", "IS 4020 (Part 8):1998", "31 Mar 1998", "Door shutters - Methods of tests Part 8", "Methods of Tests", "Indigenous"),
        ],
        generated="Generated On:\nSep 12, 2026 08:32 PM",
    )
    make_export(
        directory / "File_Published_Standards_List_2026-09-12_203626.xlsx",
        "Ministry of Chemicals and Fertilizers",
        [
            ("1", "IS 269 : 2015", None, "Ordinary Portland Cement — Specification", "-", "-"),
            ("2", "IS 7328:2020", "30 Jun 2020", "Polyethylene material for moulding and extrusion", "Product Specification", "Indigenous"),
        ],
        generated="Generated On:\nSep 12, 2026 08:36 PM",
    )
    make_export(
        directory / "File_Published_Standards_List_2026-09-12_203731.xlsx",
        "Ministry of Chemicals and Fertilizers - Department of Fertilizers (FERT)",
        [("1", "IS 14756:2024", "22 Nov 2024", "Stainless Steel Utensils — Specification ( Third Revision )", "Product Specification", "Indigenous")],
        generated="Generated On:\nSep 12, 2026 08:37 PM",
    )
    return directory


def _standard(engine, std_key):
    with engine.connect() as conn:
        return conn.execute(sa.select(schema.standard).where(schema.standard.c.std_key == std_key)).mappings().one()


def _count(engine, table):
    with engine.connect() as conn:
        return conn.execute(sa.select(sa.func.count()).select_from(table)).scalar()


def test_read_export_detects_title_generation_time_and_rows(tmp_path):
    path = tmp_path / "1.xlsx"
    make_export(path, "Total", MASTER_ROWS)
    export = read_export(path)
    assert export.title_cell == "Total"
    assert export.generated_on == datetime(2026, 9, 12, 20, 17, tzinfo=IST)
    assert len(export.rows) == 3
    assert export.rows[0].sheet_row == 3
    assert export.rows[0].designation_raw == "IS 14756:2024"
    assert export.rows[1].standard_type is None
    assert export.rows[1].degree is None
    assert export.rows[1].publish_date_raw is None


@pytest.mark.parametrize(
    "title, expected",
    [
        ("Total", "total_export"),
        (None, "untitled_export"),
        ("Ministry of Coal", "ministry_node"),
        ("Department of Atomic Energy", "ministry_node"),
        ("Ministry of Commerce and Industry - Department of Commerce", "ministry_node"),
        ("Chemical Department (CHD)", "unknown_node"),
    ],
)
def test_classify_export(tmp_path, title, expected):
    path = tmp_path / "x.xlsx"
    make_export(path, title, MASTER_ROWS[:1])
    assert classify_export(read_export(path)) == expected


def test_master_export_creates_standards_with_provenance(engine, data_dir):
    ingest_standard_exports(engine, data_dir)

    utensils = _standard(engine, "IS 14756:2024")
    assert utensils["source_id"] == "bis_std_export_total"
    assert utensils["source_locator"] == "1.xlsx!Published Standards!R3"
    assert utensils["retrieved_at"].startswith("2026-09-12T20:17")
    assert utensils["in_master_export"] is True
    assert utensils["in_second_export"] is True
    assert utensils["listing_status"] == "published_export"
    assert utensils["publication_date"] == date(2024, 11, 22)
    assert utensils["title_clean"] == "Stainless Steel Utensils — Specification (Third Revision)"
    assert utensils["revision_label"] == "Third Revision"
    assert utensils["family_key"] == "IS 14756"


def test_unspecified_values_become_null_and_are_flagged(engine, data_dir):
    ingest_standard_exports(engine, data_dir)
    cement = _standard(engine, "IS 269:2015")
    assert cement["designation_raw"] == "IS 269 : 2015"
    assert cement["standard_type"] is None
    assert cement["degree_of_equivalence"] is None
    assert cement["publication_date"] is None
    assert {"missing_publication_date", "type_unspecified", "equivalence_unspecified"} <= set(
        json.loads(cement["quality_flags"])
    )


def test_variant_spelling_is_recorded_as_alias(engine, data_dir):
    ingest_standard_exports(engine, data_dir)
    with engine.connect() as conn:
        alias = conn.execute(
            sa.select(schema.standard_alias).where(schema.standard_alias.c.raw_text == "IS 269 : 2015")
        ).mappings().first()
    assert alias is not None
    assert alias["normalized_key"] == "IS 269:2015"
    assert alias["resolution"] == "exact_version"


def test_second_export_adds_missing_designation_and_skips_duplicates(engine, data_dir):
    summary = ingest_standard_exports(engine, data_dir)
    extra = _standard(engine, "IS 4020 (Part 8):1998")
    assert extra["source_id"] == "bis_std_export_second"
    assert extra["in_master_export"] is False
    assert extra["in_second_export"] is True
    assert _standard(engine, "IS 269:2015")["in_second_export"] is False
    assert summary["duplicates_skipped"]["2.xlsx"] == 1


def test_ministry_exports_create_nodes_links_and_ministry_only_standards(engine, data_dir):
    ingest_standard_exports(engine, data_dir)

    ministry_only = _standard(engine, "IS 7328:2020")
    assert ministry_only["listing_status"] == "ministry_export_only"
    assert ministry_only["source_id"] == "bis_std_exports_ministry"

    with engine.connect() as conn:
        nodes = {
            row["export_label"]: row
            for row in conn.execute(sa.select(schema.classification_node)).mappings()
        }
        links = conn.execute(sa.select(sa.func.count()).select_from(schema.standard_classification)).scalar()

    parent = nodes["Ministry of Chemicals and Fertilizers"]
    child = nodes["Ministry of Chemicals and Fertilizers - Department of Fertilizers (FERT)"]
    assert parent["dimension"] == "ministry"
    assert parent["parent_node_id"] is None
    assert child["name"] == "Department of Fertilizers (FERT)"
    assert child["parent_node_id"] == parent["node_id"]
    assert links == 3
    assert _count(engine, schema.standard) == 5


def test_each_source_gets_a_successful_run(engine, data_dir):
    summary = ingest_standard_exports(engine, data_dir)
    assert set(summary["runs"]) == {"bis_std_export_total", "bis_std_export_second", "bis_std_exports_ministry"}
    with engine.connect() as conn:
        statuses = [row[0] for row in conn.execute(sa.select(schema.ingestion_run.c.status))]
    assert statuses == ["success", "success", "success"]


def test_reingest_is_idempotent(engine, data_dir):
    ingest_standard_exports(engine, data_dir)
    summary = ingest_standard_exports(engine, data_dir)
    for source_id, stats in summary["runs"].items():
        assert stats["inserted"] == 0, source_id
        assert stats["updated"] == 0, source_id
        assert stats["retired"] == 0, source_id
    assert _count(engine, schema.standard) == 5


def test_unclassified_export_is_skipped_and_reported(engine, data_dir):
    make_export(data_dir / "dept.xlsx", "Chemical Department (CHD)", MASTER_ROWS[:1])
    summary = ingest_standard_exports(engine, data_dir)
    assert "dept.xlsx" in summary["skipped_files"]
    with engine.connect() as conn:
        labels = {row[0] for row in conn.execute(sa.select(schema.classification_node.c.export_label))}
    assert "Chemical Department (CHD)" not in labels


def test_canonical_collision_is_rejected_not_merged(engine, tmp_path):
    directory = tmp_path / "collide"
    directory.mkdir()
    make_export(
        directory / "1.xlsx",
        "Total",
        [
            ("1", "IS 1281:2025", "01 Jan 2025", "Title A", "Product Specification", "Indigenous"),
            ("2", "IS 1281 : 2025", "01 Jan 2025", "Title B", "Product Specification", "Indigenous"),
        ],
    )
    summary = ingest_standard_exports(engine, directory)
    assert _count(engine, schema.standard) == 1
    assert summary["runs"]["bis_std_export_total"]["rejected"] == 1
    with engine.connect() as conn:
        errors = conn.execute(sa.select(schema.ingestion_run.c.errors_json)).scalar()
    assert "canonical_collision" in errors

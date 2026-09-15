"""Test helper: BIS Published Standards exports written at test time, and a fake portal that serves them."""

from pathlib import Path

import openpyxl

from manakmarg.ingest.fetch import AccessBlocked, FetchError
from manakmarg.refresh.portal import DownloadedExport, Ministry, PortalChanged

HEADERS = ["Sl#", "Standard Number", "Date of Publish", "Title", "Type of Standard", "Degree of Equivalence"]
TITLES = {"IS 13688:2020": "Packaged Pasteurized Milk — Specification ( Second Revision )"}


def rows_for(designations):
    return [
        (str(index), designation, "12 Sep 2026", TITLES.get(designation, f"Title of {designation}"), "Product Specification", "Indigenous")
        for index, designation in enumerate(designations, start=1)
    ]


def write_export(path: Path, title, rows, *, headers=HEADERS, generated="Generated On:\nSep 19, 2026 02:31 AM") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Published Standards"
    sheet["C1"] = title
    sheet.merge_cells("C1:E1")
    sheet["F1"] = generated
    for column, header in enumerate(headers, start=1):
        sheet.cell(row=2, column=column, value=header)
    for offset, row in enumerate(rows, start=3):
        for column, value in enumerate(row[: len(headers)], start=1):
            sheet.cell(row=offset, column=column, value=value)
    workbook.save(path)
    return path


def _downloaded(kind, path, ministry=None):
    content = Path(path).read_bytes()
    return DownloadedExport(kind, Path(path), "https://standards.bis.gov.in/test", "https://example.test/export", "https://example.test/file.xlsx", "0" * 64, len(content), "2026-09-19T00:00:00+00:00", ministry)


class FakePortal:
    """Serves generated exports. ``ministries`` maps a Ministry to the designations its export lists."""

    def __init__(self, total, ministries, *, total_title="Total", total_headers=HEADERS, fail=None, corrupt=None, blocked=None, list_error=None):
        self.total = total
        self.ministries = ministries
        self.total_title = total_title
        self.total_headers = total_headers
        self.fail = set(fail or ())
        self.corrupt = set(corrupt or ())
        self.blocked = set(blocked or ())
        self.list_error = list_error
        self.requested: list[str] = []

    def download_total(self, directory):
        self.requested.append("total")
        path = write_export(Path(directory) / "overall-total.xlsx", self.total_title, rows_for(self.total), headers=self.total_headers)
        return _downloaded("total", path)

    def list_ministries(self):
        if self.list_error:
            raise self.list_error
        return list(self.ministries)

    def download_ministry(self, ministry: Ministry, directory):
        self.requested.append(ministry.name)
        if ministry.ministry_id in self.blocked:
            raise AccessBlocked("https://example.test/export", "http_403", 403)
        if ministry.ministry_id in self.fail:
            raise FetchError("https://example.test/export", "HTTP 500", 500)
        path = Path(directory) / f"ministry-{ministry.ministry_id}.xlsx"
        if ministry.ministry_id in self.corrupt:
            path.write_bytes(b"PK\x03\x04 this is not a real workbook")
        else:
            write_export(path, ministry.name, rows_for(self.ministries[ministry]))
        return _downloaded("ministry", path, ministry)


__all__ = ["FakePortal", "HEADERS", "Ministry", "PortalChanged", "rows_for", "write_export"]
